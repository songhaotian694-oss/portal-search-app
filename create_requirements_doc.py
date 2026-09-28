from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parent / "中央民族大学就业分享信息检索软件需求说明书.docx"
BLUE = "173F6B"; LIGHT = "EEF3F8"; BORDER = "D9D9D9"

def set_font(run, name="Microsoft YaHei", size=10.5, bold=False, color="000000"):
    run.font.name = name; run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)

def border(cell):
    tcPr = cell._tc.get_or_add_tcPr(); borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag=OxmlElement(f"w:{edge}"); tag.set(qn("w:val"),"single"); tag.set(qn("w:sz"),"4"); tag.set(qn("w:color"),BORDER); borders.append(tag)
    tcPr.append(borders)

def cell_text(cell, text, bold=False, color="000000"):
    cell.text = ""; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(3)
    r=p.add_run(str(text)); set_font(r,size=9.5,bold=bold,color=color); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; border(cell)

def table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style="Table Grid"
    for i,h in enumerate(headers): shade(t.rows[0].cells[i],BLUE); cell_text(t.rows[0].cells[i],h,True,"FFFFFF")
    for row_no,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            if row_no % 2: shade(cells[i],LIGHT)
            cell_text(cells[i],v)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Cm(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def heading(doc, text, level=1):
    p=doc.add_paragraph(); p.style=f"Heading {level}"; p.paragraph_format.space_before=Pt(14 if level==1 else 9); p.paragraph_format.space_after=Pt(6)
    r=p.add_run(text); set_font(r,size=15 if level==1 else 12,bold=True,color="000000")

def para(doc, text, bold_lead=None):
    p=doc.add_paragraph(); p.paragraph_format.line_spacing=1.45; p.paragraph_format.space_after=Pt(5); p.paragraph_format.first_line_indent=Cm(.74)
    if bold_lead and text.startswith(bold_lead):
        r=p.add_run(bold_lead); set_font(r,bold=True); r=p.add_run(text[len(bold_lead):]); set_font(r)
    else: r=p.add_run(text); set_font(r)

def bullets(doc, items):
    for item in items:
        p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(3); r=p.add_run(item); set_font(r)

def code_block(doc, items):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(.75); p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(5)
    r=p.add_run("\n".join(items)); set_font(r,"Consolas",9.5); pPr=p._p.get_or_add_pPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),"F4F6F8");pPr.append(shd)

doc=Document(); sec=doc.sections[0]; sec.top_margin=Cm(2.2); sec.bottom_margin=Cm(2.0); sec.left_margin=Cm(2.35); sec.right_margin=Cm(2.35)
styles=doc.styles; styles["Normal"].font.name="Microsoft YaHei"; styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei"); styles["Normal"].font.size=Pt(10.5)

p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(12); r=p.add_run("中央民族大学就业分享信息检索软件需求说明书"); set_font(r,size=20,bold=True)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(22); r=p.add_run("本地运行  手工登录  本地检索与导出"); set_font(r,size=11,color="404040")
table(doc,["项目","说明"],[
    ["文档用途","明确软件范围、功能、数据、安全、测试与验收要求。"],
    ["运行边界","软件只在用户本机运行，只处理当前用户正常权限内可访问的门户内容。"],
    ["核心原则","用户手工登录；不保存账号密码；采集与处理分离；全部数据本地保存。"],
], [3.2,12.8])

heading(doc,"一 项目概述")
para(doc,"本软件用于检索中央民族大学信息门户就业信息栏目中的就业分享、选调经验等文章。系统在获得用户手工登录后的浏览器会话基础上，下载允许访问的文章及附件，在本地完成文本解析、字段提取、关键词检索、语义检索、人工核对和 Excel 导出。")
heading(doc,"1 建设目标",2)
bullets(doc,["建立本地、可持续更新的就业分享资料库。","支持按届别、年级、学历、专业、城市、单位、岗位和发布时间进行查询。","保留每项提取结果的原文依据，避免无依据推断。","在不保存账号密码、不绕过认证和不访问越权数据的条件下完成使用流程。"])
heading(doc,"2 使用对象",2)
table(doc,["对象","使用目的"],[ ["具有门户权限的学生或工作人员","查询、整理和导出就业经验资料。"],["资料核对人员","人工修订程序提取的字段并确认结果。"] ],[4,12])

heading(doc,"二 系统范围与运行方式")
para(doc,"软件为本地桌面软件，前端以 HTML、CSS 和 JavaScript 实现，使用 PyWebView 展示桌面窗口；后端使用 FastAPI，仅监听 127.0.0.1；登录与页面采集使用 Playwright 驱动可见 Chromium。")
table(doc,["类别","技术与要求"],[
["后端","Python 3.11、FastAPI、SQLite、SQLite FTS5。"],["桌面界面","PyWebView 显示本地网页界面。"],["门户采集","Playwright 打开本机可见 Chromium，用户自行认证。"],["解析","BeautifulSoup、PyMuPDF、python-docx、openpyxl、pandas、Pillow。"],["语义检索","sentence-transformers 与 BAAI/bge-small-zh-v1.5，本地缓存模型。"],["打包","PyInstaller 生成 Windows 可执行程序。"]], [3.5,12.5])

heading(doc,"三 安全与权限要求")
bullets(doc,[
"后端只绑定 127.0.0.1，不对局域网或互联网开放服务。",
"用户必须在学校官方认证网页自行输入账号、密码、验证码及完成其他验证。",
"系统不得自动填写、识别、破解或绕过验证码，不得保存明文账号和密码。",
"登录成功后仅沿用当前 Playwright BrowserContext；登录状态文件可能包含敏感 Cookie，必须保存在 data 目录并加入 .gitignore。",
"系统仅访问 allowed_domains 中的域名；外部链接只记录地址，默认不下载。",
"原始 HTML、正文、附件、数据库、索引、导出文件和日志均保存本地并加入 .gitignore。",
"日志不得记录 Cookie、密码、Authorization 请求头或其他认证敏感信息。",
"前端以纯文本方式展示正文，禁止执行文章 HTML 中的脚本。",
"下载附件前清理文件名，防止非法路径和目录穿越。"
])

heading(doc,"四 功能需求")
heading(doc,"1 首页",2); para(doc,"首页展示文章总数、已处理数量、待人工核对数量和最近同步时间，并显示同步、处理或登录过程中的状态与错误信息。")
heading(doc,"2 门户登录",2)
table(doc,["序号","需求"],[
["1","点击登录门户后启动可见 Chromium，并打开 portal_url。"],["2","用户在学校网页中自行完成认证，不向软件输入账号、密码或验证码。"],["3","用户进入就业信息栏目后，在软件中点击我已登录。"],["4","系统检查当前页面是否仍是登录页；成功后保留浏览器会话。"],["5","会话失效时暂停采集并提示重新登录；重新登录后支持继续任务。"]],[1.4,14.6])
heading(doc,"3 页面配置与检查",2)
para(doc,"真实门户结构未知时，系统必须保留 TODO，不得猜测页面选择器或接口。用户登录并进入就业栏目后可执行页面配置检查。检查时保存当前页的脱敏 HTML，分别统计选择器命中数量，显示少量脱敏文本，并明确指出空值、TODO 或无效选择器。")
code_block(doc,["portal_url", "allowed_domains", "employment_entry", "list_item_selector", "title_selector", "date_selector", "detail_link_selector", "next_button_selector", "detail_body_selector", "attachment_selector", "login_page_markers", "page_interval_seconds"])
heading(doc,"4 数据同步",2)
bullets(doc,[
"遍历就业信息列表，获取标题、发布时间、栏目、详情链接和门户文章 ID。",
"支持自动翻页、最大页数限制、测试同步、完整同步、增量同步和失败重试。",
"按详情链接或门户文章 ID 去重；按正文内容哈希判断内容是否变化。",
"每完成一页保存任务进度；逐篇保存原始 HTML 和清理后的 TXT 正文。",
"下载 PDF、Word、Excel 和图片附件，验证返回内容不是登录页或错误页。",
"保存附件名称、类型、大小、哈希和所属文章；失败记录不影响后续文章。",
"默认请求间隔为 1 至 2 秒，不进行高频请求。"
])
heading(doc,"5 本地数据处理",2)
bullets(doc,["处理阶段只能读取本地已下载内容，不得为了重新提取字段反复访问门户。","清理菜单、按钮、页脚等无关网页文字，合并正文与可解析附件文字。","PDF 使用 PyMuPDF，Word 使用 python-docx，Excel 使用 openpyxl 或 pandas。","图片和扫描 PDF 未配置 OCR 时标记待 OCR 或待人工查看，单个解析失败不得中断任务。","支持处理新增内容、重新处理全部本地内容和重建搜索索引。"])
heading(doc,"6 字段提取",2)
table(doc,["字段","提取规则","输出要求"],[
["届别","匹配 20XX届。","保存值、来源、证据原句和核对状态。"],["年级","匹配 20XX级。","保存值、来源、证据原句和核对状态。"],["学历","识别本科、硕士、博士及其常见表达。","没有依据时待人工核对。"],["专业","匹配 config/majors.txt 词表。","不得按常识补充。"],["城市","匹配 config/cities.txt 词表。","不得按常识补充。"],["单位与岗位","优先分析含选调、录用、入职、任职、单位、岗位、报到、组织部、街道、乡镇、机关的完整句子。","保存可能值及完整支撑句；证据不足时待人工核对。"]],[2.2,5.4,8.4])
heading(doc,"7 搜索与筛选",2)
para(doc,"搜索由 SQLite FTS5 关键词搜索和本地中文向量语义搜索组成。长文章按约 400 至 600 个汉字切分，文本块保存文章 ID 与位置，相邻块保留适度重叠。关键词和语义分数分别归一化后合并排序。")
code_block(doc,["综合分数 = 关键词分数 × 0.45 + 语义分数 × 0.55"])
bullets(doc,["支持按届别、学历、专业、城市和发布时间筛选。","结果显示标题、发布时间、字段信息、相关片段、匹配方式、相关程度、原文链接和人工核对状态。","模型首次使用时允许下载，后续应从本地缓存加载；模型不可用时关键词检索仍应可用。"])
heading(doc,"8 文章详情 人工核对与导出",2)
para(doc,"文章详情页展示正文摘要、提取字段、证据原句、附件和原文链接。人工核对功能应允许用户修改字段和证据，标记为已确认，并持久化到本地数据库。Excel 导出支持全部结果或当前筛选结果。")
code_block(doc,["文章标题", "发布时间", "届别", "年级", "学历", "专业", "城市", "就业单位", "岗位", "证据原句", "附件名称", "原文链接", "提取状态", "采集时间"])
para(doc,"导出文件必须额外包含采集失败记录和待人工核对记录工作表，文件保存至 data/exports 目录。")

doc.add_page_break()
heading(doc,"五 本地数据设计")
table(doc,["表名","主要内容"],[
["articles","文章 ID、门户 ID、标题、发布时间、栏目、详情链接、原始 HTML 路径、正文路径、内容哈希、状态、采集与更新时间。"],["attachments","附件 ID、文章 ID、名称、路径、类型、哈希、解析状态、提取文本路径和外部链接标识。"],["extracted_records","文章 ID、届别、年级、学历、专业、城市、单位、岗位、证据 JSON、人工核对状态。"],["sync_jobs","任务 ID、开始结束时间、采集数量、成功数量、失败数量、状态和任务说明。"],["failures","失败 URL、阶段、错误、重试次数、解决状态和创建时间。"],["article_fts 与 chunks","FTS5 检索内容；语义文本块、位置和向量数据。"]],[3.6,12.4])

heading(doc,"六 界面与接口要求")
para(doc,"界面至少包括首页、数据同步、数据处理、搜索、文章详情、人工核对、设置与页面配置检查、Excel 导出八个区域。所有耗时任务必须后台执行，界面可通过状态接口显示进度、成功数、失败数和错误信息。")
table(doc,["接口","用途"],[
["POST /api/auth/start","启动可见登录浏览器。"],["GET /api/auth/status","查询浏览器及登录状态。"],["POST /api/auth/confirm","确认用户已登录。"],["POST /api/sync/start 与 GET /api/sync/status","启动同步并读取同步进度。"],["POST /api/process/start 与 GET /api/process/status","启动本地处理并读取进度。"],["GET /api/search","关键词、语义与条件筛选搜索。"],["GET 和 PUT /api/articles/{article_id}","读取文章详情与保存人工核对。"],["GET /api/export","导出 Excel。"],["GET 和 PUT /api/settings","读取和保存页面配置。"],["POST /api/settings/test-selectors","执行页面配置检查。"]],[5.4,10.6])

heading(doc,"七 项目目录要求")
code_block(doc,["portal_search_app/", "  launcher.py  run.bat  run.ps1  requirements.txt  README.md", "  build_exe.bat  app/  config/  data/  tests/"])
para(doc,"data 目录下分别保存原始页面、正文、附件、附件文字、SQLite 数据库、搜索索引、导出文件、日志和脱敏诊断页面。")

heading(doc,"八 测试与验收要求")
heading(doc,"1 测试要求",2)
bullets(doc,["使用本地模拟 HTML 和模拟附件开展开发测试，不编造真实门户接口或真实页面选择器。","覆盖列表解析、翻页、正文、附件、字段规则、词表、去重、增量、搜索、导出、路径安全和失败隔离。","字段提取至少覆盖 20 条中文表达；运行测试后修复真实发现的问题，不得删除关键安全检查。"])
heading(doc,"2 验收标准",2)
bullets(doc,["双击 run.bat 可启动桌面软件。","可打开信息门户登录浏览器并由用户手工完成认证。","可在模拟模式下完成同步、处理、搜索、人工核对和 Excel 导出。","真实选择器配置完成后，可先采集前两页进行验证。","关闭门户后，已下载内容仍可离线处理、搜索和导出。","程序重启后，本地数据库和已下载数据仍存在。","README 包含安装、启动、登录、配置、同步、处理、搜索、导出和打包说明。"])

heading(doc,"九 真实门户启用前仍需确认的配置")
para(doc,"软件不应猜测真实就业栏目结构。用户手工登录后，需要通过页面配置检查确认就业栏目入口 employment_entry、列表项与详情页 CSS 选择器、下一页选择器、附件选择器以及登录页识别文字。确认后建议先执行测试同步，再执行完整或增量同步。")

footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=footer.add_run("中央民族大学就业分享信息检索软件需求说明书"); set_font(rr,size=8,color="666666")
doc.core_properties.title="中央民族大学就业分享信息检索软件需求说明书"; doc.core_properties.subject="软件需求说明"; doc.core_properties.author=""
doc.save(OUT); print(OUT)
