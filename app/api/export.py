from __future__ import annotations
from datetime import datetime
from fastapi import APIRouter
from fastapi.responses import FileResponse
from openpyxl import Workbook
from ..database import db
from ..settings import DIRS
router=APIRouter(prefix="/api/export",tags=["导出"])
@router.get("")
async def export(q:str="",graduation_year:str="",degree:str="",major:str="",city:str="",position:str="",date_from:str="",date_to:str=""):
    filters={'graduation_year':graduation_year,'degree':degree,'major':major,'city':city,'position':position,'date_from':date_from,'date_to':date_to}
    rows=db.search_experience_rows(q,filters,limit=None);wb=Workbook();ws=wb.active;ws.title="检索结果"
    headers=["来源文章","发布时间","姓名","届别","年级","学历","专业","城市","就业单位","岗位","OCR证据文字","原文链接","提取状态","采集时间"]
    ws.append(headers)
    for r in rows:
        ws.append([r.get("title"),r.get("published_at"),r.get("student_name"),*[None if r.get(key) in {'待人工核对','待核对'} else r.get(key) for key in ['graduation_year','grade','degree','major','city','employer','position']],r.get("evidence_text"),r.get("detail_url"),"部分信息未提取" if r.get("needs_review") else "字段较完整",r.get("collected_at")])
    partial=wb.create_sheet('信息不完整记录');partial.append(["来源文章","姓名","届别","专业","岗位","原文链接"])
    for row in rows:
        if row.get('needs_review'):
            partial.append([None if row.get(key) in {'待人工核对','待核对'} else row.get(key) for key in ['title','student_name','graduation_year','major','position','detail_url']])
    # A filtered export must not leak unrelated records through extra sheets.
    if not q and not any(filters.values()):
        failures=wb.create_sheet('采集失败记录');failures.append(["链接","阶段","错误","时间"])
        for row in db.unresolved_failures():failures.append([row[key] for key in ['url','stage','error','created_at']])
    for sheet in wb.worksheets:
        sheet.freeze_panes="A2";sheet.auto_filter.ref=sheet.dimensions
        for col in sheet.columns: sheet.column_dimensions[col[0].column_letter].width=min(45,max(12,max(len(str(x.value or "")) for x in col)+2))
    path=DIRS["exports"]/("就业检索导出_"+datetime.now().strftime("%Y%m%d_%H%M%S_%f")+".xlsx");wb.save(path)
    return FileResponse(path,filename=path.name,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
