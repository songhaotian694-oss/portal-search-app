from __future__ import annotations
from pathlib import Path
import hashlib, json, re, threading
import pymupdf
from docx import Document
from openpyxl import load_workbook

_ocr_engine=None
_ocr_lock=threading.Lock()
OCR_VERSION='layout-20260930-1'

def _ocr_rows(image)->list[dict]:
    global _ocr_engine
    if _ocr_engine is None:
        from rapidocr_onnxruntime import RapidOCR
        _ocr_engine=RapidOCR(det_limit_side_len=1280,intra_op_num_threads=2,inter_op_num_threads=1)
    result,_=_ocr_engine(image)
    return [{'box':[[float(x),float(y)] for x,y in item[0]],'text':item[1],'score':float(item[2])} for item in (result or []) if len(item)>2 and item[1]]

def _merge_rows(primary:list[dict],extra:list[dict])->list[dict]:
    from difflib import SequenceMatcher
    rows=list(primary)
    for new in extra:
        nx=sum(p[0] for p in new['box'])/4;ny=sum(p[1] for p in new['box'])/4
        height=max(p[1] for p in new['box'])-min(p[1] for p in new['box'])
        match=next((i for i,old in enumerate(rows) if abs(sum(p[1] for p in old['box'])/4-ny)<max(8,height*.55) and abs(sum(p[0] for p in old['box'])/4-nx)<max(25,height*2) and SequenceMatcher(None,old['text'],new['text']).ratio()>.55),None)
        if match is None:rows.append(new)
        elif new['score']>rows[match]['score']:rows[match]=new
    return sorted(rows,key=lambda row:(round(min(p[1] for p in row['box'])/12),min(p[0] for p in row['box'])))

def image_ocr_details(path:Path)->dict:
    """Bounded retry, original coordinates and scores cached beside each attachment."""
    import cv2,numpy as np
    digest=hashlib.sha256(path.read_bytes()).hexdigest()
    cache=path.with_suffix(path.suffix+'.ocr.json')
    if cache.is_file():
        try:
            saved=json.loads(cache.read_text(encoding='utf-8'))
            if saved.get('version')==OCR_VERSION and saved.get('sha256')==digest:return saved
        except (ValueError,OSError):pass
    image=cv2.imdecode(np.frombuffer(path.read_bytes(),dtype=np.uint8),cv2.IMREAD_COLOR)
    if image is None:raise ValueError('无法解码图片')
    with _ocr_lock:
        rows=_ocr_rows(image);passes=1
        text='\n'.join(row['text'] for row in rows)
        retry=not rows or any(row['score']<.88 and len(row['text'])>=4 and re.search(r'20\d{2}|专业|学士|硕士|博士|本科|公务员|科员|省|县|市|局|政府',row['text']) for row in rows) or bool(re.search(r'(?:不定|试用期公|专)\n',text))
        if retry and image.shape[0]>600:
            height=image.shape[0]
            for top,bottom in [(0,int(height*.58)),(int(height*.42),height)]:
                extra=_ocr_rows(image[top:bottom])
                for row in extra:
                    row['box']=[[x,y+top] for x,y in row['box']]
                rows=_merge_rows(rows,extra);passes+=1
    saved={'version':OCR_VERSION,'sha256':digest,'passes':passes,'lines':rows,'text':'\n'.join(row['text'] for row in rows)}
    temporary=cache.with_suffix(cache.suffix+'.tmp')
    temporary.write_text(json.dumps(saved,ensure_ascii=False),encoding='utf-8');temporary.replace(cache)
    return saved

def image_ocr(path:Path)->str:
    return image_ocr_details(path)['text']

def extract_attachment_text(path: Path) -> tuple[str, str]:
    """返回附件文字和状态；就业分享海报在本机使用中文OCR识别。"""
    try:
        ext=path.suffix.lower()
        if ext==".pdf":
            pages=[];used_ocr=False
            with pymupdf.open(path) as doc:
                for page in doc:
                    text=page.get_text()
                    if len(text.strip())<20:
                        with _ocr_lock:
                            scale=min(2,2400/max(page.rect.width,page.rect.height))
                            pix=page.get_pixmap(matrix=pymupdf.Matrix(scale,scale),alpha=False)
                            text='\n'.join(row['text'] for row in _ocr_rows(pix.tobytes('png')))
                        used_ocr=True
                    pages.append(text)
            text='\n'.join(pages)
            return (text, ('ocr_parsed' if used_ocr else 'parsed') if text.strip() else 'ocr_empty')
        if ext==".docx":
            doc=Document(path); rows=[]
            for p in doc.paragraphs: rows.append(p.text)
            for table in doc.tables:
                rows.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            return "\n".join(rows),"parsed"
        if ext in {".xlsx",".xlsm"}:
            wb=load_workbook(path,data_only=True,read_only=True); rows=[]
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True): rows.append(" | ".join(str(x) for x in row if x is not None))
            return "\n".join(rows),"parsed"
        if ext in {".png",".jpg",".jpeg",".bmp",".tiff",".tif",".webp"}:
            text=image_ocr(path)
            return text,("ocr_parsed" if text.strip() else "ocr_empty")
        return "","unsupported"
    except Exception as exc: return f"[附件解析失败：{type(exc).__name__}]","failed"
