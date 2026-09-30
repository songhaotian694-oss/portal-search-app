"""Isolated portable-package smoke check, without school authentication."""
from __future__ import annotations
import argparse
import asyncio
import json
import sys
import tempfile
from pathlib import Path


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument('--self-check', action='store_true')
    parser.add_argument('--self-check-report', required=True)
    options = parser.parse_args()
    report = {'ok': False, 'checks': []}
    try:
        with tempfile.TemporaryDirectory(prefix='portal-check-') as directory:
            from . import settings
            # Set the temporary root before importing modules that create a DB.
            root = Path(directory)
            settings.ROOT = root
            settings.CONFIG_DIR = root / 'config'
            settings.DATA_DIR = root / 'data'
            settings.DIRS = {key: settings.DATA_DIR / key for key in settings.DIRS}
            from .database import db
            from .processing.attachment_parser import extract_attachment_text, _ocr_rows
            from PIL import Image, ImageDraw, ImageFont
            font_path = Path('C:/Windows/Fonts/msyh.ttc')
            if not font_path.is_file(): font_path = Path('C:/Windows/Fonts/arial.ttf')
            image = Image.new('RGB', (850, 120), 'white')
            ImageDraw.Draw(image).text((25, 25), '2024届计算机硕士', fill='black', font=ImageFont.truetype(str(font_path), 44))
            recognized = ''.join(row['text'] for row in _ocr_rows(image))
            assert '2024' in recognized, 'OCR inference did not recognize the synthetic text'
            report['checks'].append('ocr_inference')

            from docx import Document
            document = Document(); document.add_paragraph('Portable document check')
            word = root / 'check.docx'; document.save(word)
            text, status = extract_attachment_text(word)
            assert 'Portable document check' in text and status == 'parsed', f'Word parse status: {status}'
            import pymupdf
            pdf = pymupdf.open(); page = pdf.new_page(); page.insert_text((50, 50), 'Portable PDF document verification')
            pdf_path = root / 'check.pdf'; pdf.save(pdf_path); pdf.close()
            text, status = extract_attachment_text(pdf_path)
            assert 'Portable PDF document verification' in text and status == 'parsed', f'PDF parse status: {status}'
            report['checks'].append('word_pdf')

            aid, _ = db.upsert_article({'title': '选调经验分享', 'detail_url': 'https://example.test/check'})
            db.replace_experience_records(aid, [{'student_name': '测试同学', 'graduation_year': '2024届', 'city': '北京市', 'degree': '硕士研究生', 'major': '计算机科学与技术'}])
            from .api import search, export
            result = asyncio.run(search.search(q='2024届北京计科硕士', limit=100))
            assert len(result['results']) == 1
            response = asyncio.run(export.export(q='2024届北京计科硕士'))
            from openpyxl import load_workbook
            assert load_workbook(response.path)['检索结果']['C2'].value == '测试同学'
            report['checks'].append('search_excel')

            from playwright.async_api import async_playwright
            async def driver_check():
                async with async_playwright() as playwright:
                    assert playwright.chromium.name == 'chromium'
            asyncio.run(driver_check())
            report['checks'].append('playwright_driver')
            from webview.platforms import winforms
            assert winforms.BrowserView is not None
            report['checks'].append('webview2_runtime')
            assert not any(name in sys.modules for name in ('torch', 'transformers', 'scipy', 'pandas'))
            report['checks'].append('no_unused_ml_stack')
            report['ok'] = True
    except Exception as exc:
        report['error'] = f'{type(exc).__name__}: {exc}'
    Path(options.self_check_report).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    return 0 if report['ok'] else 1
