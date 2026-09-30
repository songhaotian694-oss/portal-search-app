import asyncio
from unittest.mock import AsyncMock

import pytest
from openpyxl import load_workbook

from app.api import sync, process, export as export_api
from app.database import Database
from app.models import ProcessRequest, SyncRequest


@pytest.fixture
def local_db(tmp_path, monkeypatch):
    database = Database(tmp_path / 'recovery.db')
    for module in (sync, process, export_api):
        monkeypatch.setattr(module, 'db', database)
    monkeypatch.setattr(sync, 'task', None)
    monkeypatch.setattr(process, 'task', None)
    monkeypatch.setattr(sync, '_ensured', False)
    monkeypatch.setattr(sync, 'STATE', dict(status='idle', current=0, total=0, success=0, failed=0, message=''))
    monkeypatch.setattr(process, 'STATE', dict(status='idle', current=0, total=0, failed=0, message=''))
    directories = {}
    for key in ('raw_json', 'extracted_text', 'exports'):
        directories[key] = tmp_path / key
        directories[key].mkdir()
    for module in (sync, process, export_api):
        monkeypatch.setattr(module, 'DIRS', directories)
    return database


def api_config():
    return {'data_source': 'api', 'employment_entry': 'https://portal.test/notices',
            'allowed_domains': ['portal.test'], 'api_keywords': ['选调'],
            'api_detail_url_template': 'https://portal.test/print?id={notice_id}',
            'page_interval_seconds': 0}


def notice(number):
    return {'notice_id': number, 'notice_title': '选调经验分享', 'notice_content': '<p>正文</p>'}


def test_api_download_resumes_pending_record_and_next_page_after_restart(local_db, monkeypatch):
    cfg = api_config()
    req = SyncRequest(mode='full', max_pages=1000)
    local_db.save_checkpoint('sync', {'request': req.model_dump()})
    monkeypatch.setattr(sync, 'capture_first_page', AsyncMock(return_value=({}, {})))
    monkeypatch.setattr(sync, 'save_raw_page', lambda *args: None)
    pages = []
    downloaded = []
    interrupted = True

    async def fetch(page, template, offset, keyword):
        pages.append(offset)
        return {'datas': {'tables': [notice(1), notice(2)] if offset == 0 else []}}

    async def save(page, item, row, config):
        nonlocal interrupted
        downloaded.append(row['notice_id'])
        if row['notice_id'] == 2 and interrupted:
            interrupted = False
            raise asyncio.CancelledError()
        return item, []

    monkeypatch.setattr(sync, 'fetch_page', fetch)
    monkeypatch.setattr(sync, 'save_api_article', save)
    with pytest.raises(asyncio.CancelledError):
        asyncio.run(sync.api_portal_sync(req, 1, object(), cfg, cfg['employment_entry']))
    assert local_db.checkpoint('sync')['page_offset'] == 1
    assert [work['payload']['row']['notice_id'] for work in local_db.work_items('sync', ('pending',))] == [2]

    # A new connection simulates process restart; no in-memory task is reused.
    reloaded = Database(local_db.path)
    monkeypatch.setattr(sync, 'db', reloaded)
    monkeypatch.setattr(sync, 'load_config', lambda: cfg)
    monkeypatch.setattr(sync.browser_session, 'is_logged_in', AsyncMock(return_value=(True, '')))
    monkeypatch.setattr(sync.browser_session, 'page', object())
    asyncio.run(sync.portal_sync(req, 2,))
    assert pages == [0, 1]
    assert downloaded == [1, 2, 2]
    assert not reloaded.work_items('sync')
    assert reloaded.stats()['articles'] == 2


def test_login_automatically_resumes_even_when_partial_database_is_not_empty(local_db, monkeypatch):
    local_db.upsert_article({'title': '选调经验分享', 'detail_url': 'https://portal.test/one'})
    local_db.save_checkpoint('sync', {'request': SyncRequest(mode='full').model_dump(), 'page_offset': 7})
    monkeypatch.setattr(sync.browser_session, 'is_logged_in', AsyncMock(return_value=(True, '')))
    start = AsyncMock()
    monkeypatch.setattr(sync, 'start', start)
    result = asyncio.run(sync.ensure())
    assert result == {'started': True, 'stage': 'sync'}
    assert start.await_args.args[0].mode == 'resume'


def test_retry_targets_saved_failure_only_and_then_recognizes_downloaded_article(local_db, monkeypatch):
    item = {'title': '后面页码的选调经验分享', 'detail_url': 'https://portal.test/later'}
    local_db.queue_work('sync', item['detail_url'], {'item': item, 'row': notice(999), 'source': 'api'})
    local_db.finish_work('sync', item['detail_url'], 'network interrupted')
    local_db.add_failure(None, item['detail_url'], 'api_download', 'network interrupted')
    save = AsyncMock(return_value=(item, []))
    recognize = AsyncMock()
    monkeypatch.setattr(sync, 'save_api_article', save)
    monkeypatch.setattr(process, 'run', recognize)
    # No list-page requests are needed to retry this record.
    forbidden = AsyncMock(side_effect=AssertionError('must not enumerate pages'))
    monkeypatch.setattr(sync, 'capture_first_page', forbidden)
    asyncio.run(sync.retry_failed_items(1, object(), api_config()))
    assert save.await_count == 1
    forbidden.assert_not_awaited()
    assert recognize.await_args.args[0].article_ids == [1]
    assert local_db.unresolved_failures() == []
    with local_db.connection() as connection:
        failure = connection.execute('SELECT retry_count,resolved FROM failures').fetchone()
    assert tuple(failure) == (1, 1)


def test_failed_retry_updates_attempt_count_without_duplicate_failures(local_db, monkeypatch):
    item = {'title': '选调经验分享', 'detail_url': 'https://portal.test/fail'}
    local_db.queue_work('sync', item['detail_url'], {'item': item, 'row': notice(1), 'source': 'api'})
    local_db.finish_work('sync', item['detail_url'], 'failed')
    local_db.add_failure(None, item['detail_url'], 'api_download', 'failed')
    monkeypatch.setattr(sync, 'save_api_article', AsyncMock(side_effect=RuntimeError('still offline')))
    asyncio.run(sync.retry_failed_items(1, object(), api_config()))
    asyncio.run(sync.retry_failed_items(2, object(), api_config()))
    failures = local_db.unresolved_failures()
    assert len(failures) == 1 and failures[0]['retry_count'] == 2
    assert local_db.work_items('sync')[0]['status'] == 'failed'


def test_ocr_restart_processes_remaining_articles_only(local_db, tmp_path, monkeypatch):
    ids = []
    for index in (1, 2):
        text = tmp_path / f'{index}.txt'
        text.write_text('张同学\n2024届选调生\n法学本科生', encoding='utf-8')
        aid, _ = local_db.upsert_article({'title': '选调经验分享', 'detail_url': f'https://portal.test/{index}', 'text_path': str(text)})
        ids.append(aid)
    replace = local_db.replace_experience_records
    seen = []

    def save(article_id, records):
        seen.append(article_id)
        if article_id == ids[1] and seen.count(article_id) == 1:
            raise asyncio.CancelledError()
        return replace(article_id, records)

    monkeypatch.setattr(local_db, 'replace_experience_records', save)
    with pytest.raises(asyncio.CancelledError):
        asyncio.run(process.run(ProcessRequest(all_local=True, build_index=False)))
    assert local_db.checkpoint('process')['remaining_ids'] == [ids[1]]
    asyncio.run(process.run(ProcessRequest(build_index=False)))
    assert seen == [ids[0], ids[1], ids[1]]
    assert local_db.checkpoint('process') is None
    assert len(local_db.experience_rows()) == 2


def test_export_filters_main_and_partial_sheets_including_end_date(local_db):
    for name, city, position, date in [('甲', '北京', '科员', '2025-12-31 20:00:00'), ('乙', '上海', '科员', '2025-01-01'), ('丙', '北京', '主任', '2025-01-01')]:
        aid, _ = local_db.upsert_article({'title': '选调经验分享', 'detail_url': f'https://portal.test/{name}', 'published_at': date})
        local_db.replace_experience_records(aid, [{'student_name': name, 'city': city, 'position': position, 'degree': '本科生', 'graduation_year': '2024届', 'needs_review': True}])
    result = asyncio.run(export_api.export(q='甲', city='北京', position='科员', graduation_year='2024', degree='本科', date_from='2025-01-01', date_to='2025-12-31'))
    book = load_workbook(result.path)
    assert list(book['检索结果'].values)[1][2] == '甲'
    assert book['检索结果'].max_row == 2
    assert book['信息不完整记录'].max_row == 2
    assert '采集失败记录' not in book.sheetnames
    empty = load_workbook(asyncio.run(export_api.export(city='不存在')).path)
    assert empty['检索结果'].max_row == 1
    assert empty['信息不完整记录'].max_row == 1
    all_rows = load_workbook(asyncio.run(export_api.export()).path)
    assert all_rows['检索结果'].max_row == 4
    assert '采集失败记录' in all_rows.sheetnames


def test_filtered_database_query_can_export_without_default_limit(local_db):
    aid, _ = local_db.upsert_article({'title': '选调经验分享', 'detail_url': 'https://portal.test/many'})
    local_db.replace_experience_records(aid, [{'student_name': str(number), 'city': '北京'} for number in range(75)])
    assert len(local_db.search_experience_rows('', {'city': '北京'}, limit=None)) == 75
    assert len(local_db.search_experience_rows('', {'city': '北京'})) == 50


def test_ocr_retry_only_processes_the_failed_attachment_article(local_db, tmp_path, monkeypatch):
    ids = []
    for number in (1, 2):
        aid, _ = local_db.upsert_article({'title': '选调经验分享', 'detail_url': f'https://portal.test/{number}'})
        local_db.replace_experience_records(aid, [{'student_name': f'原记录{number}'}])
        local_db.add_attachment(aid, {'file_path': str(tmp_path / f'{number}.jpg'), 'file_hash': str(number)})
        ids.append(aid)
    attachment = local_db.get_article(ids[1])['attachments'][0]
    local_db.add_failure(None, f"attachment:{attachment['id']}", 'auto_recognition', 'failed')
    calls = []

    def recognize(path):
        calls.append(path.name)
        return '张同学\n2024届选调生\n法学本科生', 'ocr_parsed'

    monkeypatch.setattr(process, 'extract_attachment_text', recognize)
    asyncio.run(sync.retry_failed_items(1, object(), api_config()))
    assert calls == ['2.jpg']
    assert local_db.experience_rows()[0]['student_name'] == '原记录1'
    assert not local_db.unresolved_failures()
    with local_db.connection() as connection:
        assert connection.execute('SELECT retry_count FROM failures').fetchone()[0] == 1


def test_sync_page_work_and_cursor_commit_together(local_db):
    local_db.save_checkpoint('sync', {'page_offset': 3})
    payload = {'source': 'api', 'item': {'detail_url': 'https://portal.test/valid'}, 'row': {}}
    invalid = {'source': 'api', 'item': {'detail_url': 'https://portal.test/invalid'}, 'row': object()}
    with pytest.raises(TypeError):
        local_db.stage_sync_page([payload, invalid], {'page_offset': 4}, {})
    assert local_db.checkpoint('sync')['page_offset'] == 3
    assert local_db.work_items('sync') == []


def test_ocr_checkpoint_exists_before_background_task_starts(local_db):
    aid, _ = local_db.upsert_article({'title': '选调经验分享', 'detail_url': 'https://portal.test/handoff'})

    async def check():
        await process.start(ProcessRequest(build_index=False))
        assert local_db.checkpoint('process')['remaining_ids'] == [aid]
        process.task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await process.task

    asyncio.run(check())
    assert Database(local_db.path).checkpoint('process')['remaining_ids'] == [aid]
