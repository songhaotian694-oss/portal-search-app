import asyncio
import pytest
from openpyxl import load_workbook
from app.database import Database


@pytest.fixture
def database(tmp_path):
    db = Database(tmp_path / 'query.db')
    aid, _ = db.upsert_article({'title': '北京选调经验分享', 'detail_url': 'https://example.test/one', 'published_at': '2024-12-31T17:00:00'})
    db.replace_experience_records(aid, [
        {'student_name': '张三', 'graduation_year': '2024届', 'grade': '2020级', 'city': '北京市', 'degree': '硕士研究生', 'major': '计算机科学与技术', 'employer': '北京市财政局', 'position': '科员'},
        {'student_name': '李四', 'graduation_year': '2023届', 'grade': '2019级', 'city': '上海市', 'degree': '本科生', 'major': '法学', 'employer': '上海市法院', 'position': '公务员'},
        {'student_name': '王五', 'graduation_year': '2025届', 'grade': '2021级', 'city': '上海市', 'degree': '博士', 'major': '计算机科学', 'employer': '上海市财政局', 'position': '科员'},
        {'student_name': '赵六', 'graduation_year': '2024届', 'grade': '2020级', 'city': '北京市', 'degree': '待人工核对', 'major': '软件工程', 'employer': '某单位', 'position': '科员'},
        {'student_name': '周七', 'graduation_year': '2024届', 'grade': '2020级', 'city': '邯郸市', 'degree': '学士学位', 'major': '法律', 'employer': '邯郸市法院', 'position': '公务员'},
    ])
    return db


def names(db, query, filters=None):
    return {row['student_name'] for row in db.search_experience_rows(query, filters)}


@pytest.mark.parametrize(('query', 'expected'), [
    ('帮我找一下在北京市工作的2024届计科硕士', {'张三'}),
    ('北京或上海，排除本科', {'张三', '王五'}),
    ('北京或上海，不看本科', {'张三', '王五'}),
    ('不在北京或上海工作的本科生', {'周七'}),
    ('2023至2025届，法学或计算机', {'张三', '李四', '王五', '周七'}),
    ('2020级北京', {'张三', '赵六'}),
    ('2024年北京', {'张三', '赵六'}),
    ('2024年毕业的北京同学', {'张三', '赵六'}),
    ('张三财政局', {'张三'}),
    ('单位：法院 姓名：李四', {'李四'}),
    ('法律学士', {'李四', '周七'}),
    ('2024届邯郸法律', {'周七'}),
    ('研究生 财政局', {'张三', '王五'}),
    ('财政局或法院', {'张三', '李四', '王五', '周七'}),
    ('北京 排除"财政局"', {'赵六'}),
    ('CS硕士', {'张三'}),
    ('２０２４届 计科 硕士', {'张三'}),
    ('2025至2023届', set()),
])
def test_natural_query(database, query, expected):
    assert names(database, query) == expected


def test_field_override_and_publication_year_are_independent(database):
    assert names(database, '北京', {'city': '上海'}) == {'李四', '王五'}
    assert names(database, '2025届') == {'王五'}
    assert names(database, '2025年') == set()
    assert names(database, '北京') == {'张三', '赵六'}  # shared article title is not a person's city
    assert names(database, '2020届') == set()  # grade cannot masquerade as graduation year
    assert names(database, '2024年', {'date_from': '2025-01-01', 'date_to': '2025-12-31'}) == set()


def test_literals_are_not_sql_wildcards(database):
    for query in ('%', '_', '"%"', '"_"', "' OR 1=1 --"):
        assert names(database, query) == set()


def test_significant_chinese_words_are_not_deleted(database):
    from app.processing.chinese_query import parse_query
    assert parse_query('信息中心').terms == [['信息中心']]
    assert parse_query('在信息中心工作的').terms == [['信息中心']]


def test_export_and_api_use_identical_natural_query(database, tmp_path, monkeypatch):
    from app.api import search, export
    monkeypatch.setattr(search, 'db', database)
    monkeypatch.setattr(export, 'db', database)
    monkeypatch.setitem(export.DIRS, 'exports', tmp_path)
    query = '北京或上海，排除本科'
    result = asyncio.run(search.search(q=query, limit=100))
    response = asyncio.run(export.export(q=query))
    wb = load_workbook(response.path)
    assert {row[2] for row in wb['检索结果'].iter_rows(min_row=2, values_only=True)} == {row['student_name'] for row in result['results']} == {'张三', '王五'}


def test_chunk_index_has_no_model_dependency(database, monkeypatch):
    from app.processing import search_index
    monkeypatch.setattr(search_index, 'db', database)
    search_index.search_index.index_article(1, '计算机专业同学分享选调经历。' * 60)
    with database.connection() as connection:
        rows = connection.execute('SELECT text,vector_json FROM chunks').fetchall()
    assert len(rows) > 1 and all(row['vector_json'] is None for row in rows)
