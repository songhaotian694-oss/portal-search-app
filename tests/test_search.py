from app.database import Database,fts_terms
from app.processing.text_chunker import chunk_text
def test_fts_and_filters(tmp_path):
    d=Database(tmp_path/"x.db");aid,_=d.upsert_article({"title":"北京选调经验","detail_url":"https://x/a","content_hash":"x","published_at":"2024-01-01"})
    d.update_fts(aid,"北京选调经验","软件工程专业同学被组织部录用到北京工作")
    d.save_record(aid,{"graduation_year":"2024届","grade":"2020级","degree":"本科","major":"软件工程","city":"北京","employer":"组织部","position":"公务员","evidence":{},"needs_review":False})
    with d.connection() as c:assert c.execute("SELECT count(*) FROM article_fts WHERE article_fts MATCH ?",(fts_terms("北京"),)).fetchone()[0]==1
    assert len(d.all_rows({"city":"北京"}))==1 and not d.all_rows({"city":"上海"})
def test_chunk_overlap():
    parts=chunk_text("甲"*560+"。"+"乙"*300);assert len(parts)>=2 and parts[0][-1] in "甲。" and len(parts[1])>0

def test_excel_export(tmp_path,monkeypatch):
    import asyncio
    from app.api import export as export_api
    from app.settings import DIRS
    d=Database(tmp_path/"x.db");aid,_=d.upsert_article({"title":"导出测试","detail_url":"https://x/export","content_hash":"x"})
    d.save_record(aid,{"graduation_year":"2024届","grade":"2020级","degree":"本科","major":"软件工程","city":"北京","employer":"待人工核对","position":"待人工核对","evidence":{},"needs_review":True})
    d.replace_experience_records(aid,[{"student_name":"张同学","graduation_year":"待人工核对","grade":"2020级","degree":"本科","major":"软件工程","city":"北京","employer":"某单位","position":"待人工核对","needs_review":True}])
    monkeypatch.setattr(export_api,"db",d); response=asyncio.run(export_api.export())
    from openpyxl import load_workbook
    wb=load_workbook(response.path)
    assert "检索结果" in wb.sheetnames and "信息不完整记录" in wb.sheetnames
    assert wb["信息不完整记录"]["B2"].value=="张同学"
    assert wb['检索结果']['D2'].value is None
